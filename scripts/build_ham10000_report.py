from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK = ROOT / "kaggle_sin_notebook706b5506c8.ipynb"
OUT_DOCX = ROOT / "HAM10000_ConvNeXt_최종보고서.docx"
CHART_DIR = ROOT / "report_assets"

FONT_KR = "Apple SD Gothic Neo"
FONT_LATIN = "Arial"
HEADING_BLUE = RGBColor(26, 88, 138)
HEADING_LIGHT = RGBColor(39, 122, 184)
MUTED = RGBColor(91, 99, 112)
ACCENT = "#277AB8"
DARK = "#111827"
TEAL = "#2A9D8F"
AMBER = "#E9A23B"
CORAL = "#D66A5A"
SOFT_BLUE = "#EAF4FF"
SOFT_GRAY = "#F5F7FA"


def require_notebook() -> None:
    if not NOTEBOOK.exists():
        raise FileNotFoundError(f"Notebook not found: {NOTEBOOK}")
    nb = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    if len(nb.get("cells", [])) < 70:
        raise RuntimeError("Notebook appears incomplete; expected latest experiment cells.")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None) -> None:
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
    normal.font.size = Pt(10.4)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)

    for style_name, size, color, before, after in [
        ("Heading 1", 17, HEADING_LIGHT, 18, 9),
        ("Heading 2", 13.5, HEADING_BLUE, 13, 6),
        ("Heading 3", 11.5, HEADING_BLUE, 9, 4),
    ]:
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("HAM10000 ConvNeXt 최종 보고서")
    set_run_font(r, 8.2, False, MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("최종 보고서")
    set_run_font(r, 8.2, False, MUTED)


def add_para(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(7)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_after = Pt(3)


def _is_numeric_cell(value: str) -> bool:
    s = str(value).strip()
    allowed = set("0123456789.+-→%≈/ \n")
    return bool(s) and all(ch in allowed for ch in s)


def _is_text_column(header: str) -> bool:
    text_headers = {
        "실험",
        "목적",
        "결과",
        "결론",
        "해석",
        "핵심 해석",
        "채택 여부",
        "채택/미채택 근거",
        "보고서에서 중요한 이유",
        "핵심 변화",
        "대표 실험",
        "범주",
    }
    return header in text_headers or "해석" in header or "근거" in header or "이유" in header


def add_table(
    doc: Document,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
    body_size: float = 8.0,
    header_size: float = 8.2,
    margin_y: int = 80,
    margin_x: int = 120,
):
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    r = caption.add_run(title)
    set_run_font(r, 9.2, True, HEADING_BLUE)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, header_size, True)
        set_cell_shading(hdr[i], "EAF4FF")
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(hdr[i], top=margin_y, bottom=margin_y, start=margin_x, end=margin_x)
        if widths:
            set_cell_width(hdr[i], widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            header = headers[i]
            if _is_text_column(header) or (len(str(val)) >= 18 and not _is_numeric_cell(str(val))):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            set_run_font(r, body_size)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i], top=margin_y, bottom=margin_y, start=margin_x, end=margin_x)
            if widths:
                set_cell_width(cells[i], widths[i])

    doc.add_paragraph()
    return table


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F5F9FF")
    set_cell_margins(cell, top=140, bottom=140, start=170, end=170)
    set_cell_width(cell, 16.3)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    set_run_font(r, 9.6, False, RGBColor(31, 41, 55))
    doc.add_paragraph()


def setup_chart_font() -> None:
    candidates = [
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/Library/Fonts/NanumGothic.ttc",
        "/System/Library/Fonts/ヒラギノ角ゴシック W4.ttc",
    ]
    for c in candidates:
        if Path(c).exists():
            font_manager.fontManager.addfont(c)
            prop = font_manager.FontProperties(fname=c)
            plt.rcParams["font.family"] = prop.get_name()
            break
    plt.rcParams["axes.unicode_minus"] = False


def save_bar_chart(path: Path, title: str, labels: list[str], values: list[float], ylabel: str, color: str = ACCENT, ylim=None) -> None:
    setup_chart_font()
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=180)
    chart_colors = [ACCENT, TEAL, AMBER, CORAL, "#7C3AED", "#475569"]
    bar_colors = chart_colors[: len(labels)] if color == ACCENT else color
    bars = ax.bar(labels, values, color=bar_colors, alpha=0.92)
    ax.set_title(title, fontsize=13, fontweight="bold", color=DARK, pad=12)
    ax.set_ylabel(ylabel, fontsize=9)
    if ylim:
        ax.set_ylim(*ylim)
    ax.grid(axis="y", alpha=0.22)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="x", labelrotation=0, labelsize=8)
    ax.tick_params(axis="y", labelsize=8)
    for b, v in zip(bars, values):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height(), f"{v:.4f}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_line_chart(path: Path, title: str, xs: list[float], series: dict[str, list[float]], ylabel: str) -> None:
    setup_chart_font()
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=180)
    colors = [ACCENT, CORAL, TEAL, AMBER]
    for (name, ys), color in zip(series.items(), colors):
        ax.plot(xs, ys, marker="o", linewidth=2, markersize=4, label=name, color=color)
    ax.set_title(title, fontsize=13, fontweight="bold", color=DARK, pad=12)
    ax.set_xlabel("OVR blending alpha", fontsize=9)
    ax.set_ylabel(ylabel, fontsize=9)
    ax.grid(alpha=0.22)
    ax.legend(fontsize=8, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_delta_chart(path: Path) -> None:
    setup_chart_font()
    labels = ["mel recall", "mel F1", "bcc recall", "bcc F1", "df recall", "df F1"]
    before = [0.4933, 0.5946, 0.8155, 0.7467, 0.5217, 0.6154]
    after = [0.6009, 0.6247, 0.9029, 0.7718, 0.6522, 0.7143]
    fig, ax = plt.subplots(figsize=(7.2, 3.8), dpi=180)
    x = range(len(labels))
    ax.bar([i - 0.18 for i in x], before, width=0.36, label="before", color="#A9B4C2")
    ax.bar([i + 0.18 for i in x], after, width=0.36, label="after", color=TEAL)
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_ylim(0.35, 0.95)
    ax.set_title("Threshold 조정 전후: recall/F1 변화", fontsize=13, fontweight="bold", color=DARK, pad=12)
    ax.grid(axis="y", alpha=0.22)
    ax.legend(fontsize=8, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def build_charts() -> dict[str, Path]:
    CHART_DIR.mkdir(exist_ok=True)
    charts = {
        "backbone": CHART_DIR / "chart_backbone_auroc.png",
        "imbalance": CHART_DIR / "chart_imbalance_auroc.png",
        "calibration": CHART_DIR / "chart_calibration.png",
        "threshold": CHART_DIR / "chart_threshold.png",
        "ovr": CHART_DIR / "chart_ovr_alpha.png",
    }
    save_bar_chart(
        charts["backbone"],
        "Backbone별 best validation Macro-AUROC",
        ["ResNet18", "EffNet-B0", "ConvNeXt-Tiny"],
        [0.898899, 0.894029, 0.940512],
        "Macro-AUROC",
        ylim=(0.86, 0.95),
    )
    save_bar_chart(
        charts["imbalance"],
        "불균형 대응 실험 Macro-AUROC 비교",
        ["Baseline", "Focal", "WRS", "CB Loss"],
        [0.966819, 0.9098, 0.9505, 0.9457],
        "Macro-AUROC",
        color=[ACCENT, CORAL, TEAL, AMBER],
        ylim=(0.88, 0.98),
    )
    setup_chart_font()
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=180)
    metrics = ["NLL", "ECE"]
    before = [0.464353, 0.044550]
    after = [0.439376, 0.019226]
    x = range(len(metrics))
    ax.bar([i - 0.18 for i in x], before, width=0.36, label="before", color="#A9B4C2")
    ax.bar([i + 0.18 for i in x], after, width=0.36, label="after", color=TEAL)
    ax.set_xticks(list(x))
    ax.set_xticklabels(metrics, fontsize=9)
    ax.set_title("Temperature Scaling 전후 확률 품질 변화", fontsize=13, fontweight="bold", color=DARK, pad=12)
    ax.grid(axis="y", alpha=0.22)
    ax.legend(fontsize=8, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    for i, (b, a) in enumerate(zip(before, after)):
        ax.text(i - 0.18, b, f"{b:.3f}", ha="center", va="bottom", fontsize=8)
        ax.text(i + 0.18, a, f"{a:.3f}", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    fig.savefig(charts["calibration"], bbox_inches="tight")
    plt.close(fig)
    save_delta_chart(charts["threshold"])
    save_line_chart(
        charts["ovr"],
        "mel precision-recall trade-off by OVR blending",
        [0.00, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.35, 0.40, 0.45, 0.50, 0.55, 0.60, 0.65, 0.70, 0.75, 0.80],
        {
            "mel recall": [0.493274, 0.493274, 0.493274, 0.511211, 0.515695, 0.515695, 0.524664, 0.538117, 0.542601, 0.556054, 0.569507, 0.578475, 0.591928, 0.605381, 0.614350, 0.645740, 0.668161],
            "mel precision": [0.748299, 0.738255, 0.728477, 0.730769, 0.732484, 0.732484, 0.726708, 0.722892, 0.703488, 0.696629, 0.686486, 0.668394, 0.653465, 0.627907, 0.595652, 0.573705, 0.564394],
            "mel F1": [0.594595, 0.591398, 0.588235, 0.601583, 0.605263, 0.605263, 0.609375, 0.616967, 0.612658, 0.618454, 0.622549, 0.620192, 0.621176, 0.616438, 0.604857, 0.607595, 0.611910],
        },
        "score",
    )
    return charts


def add_chart(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(caption)
    set_run_font(r, 9.2, True, HEADING_BLUE)
    doc.add_picture(str(path), width=Inches(6.5))
    doc.add_paragraph()


def build_report() -> None:
    require_notebook()
    charts = build_charts()
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("HAM10000 피부 병변 분류 최종 보고서")
    set_run_font(r, 22, True, RGBColor(18, 30, 48))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = meta.add_run(
        "강사님 '더 해볼 것' 구현 결과 + 이후 추가 실험 반영판\n"
        f"기준 노트북: {NOTEBOOK.name} / 작성일: {date.today().isoformat()}"
    )
    set_run_font(r, 9.8, False, MUTED)

    add_callout(
        doc,
        "이 보고서는 기존 HAM10000_실험결과_리포트_개정판.pdf에 담긴 '더 해볼 것' 구현 결과를 출발점으로 삼고, "
        "그 이후 노트북에서 추가된 Focal Loss, WeightedRandomSampler, Class-Balanced Loss, OVR Expert, Ensemble, Temperature Scaling, Threshold Optimization까지 이어 붙여 정리한다. "
        "핵심은 실험을 많이 했다는 사실이 아니라, 처음의 판단 기준이 어떻게 바뀌었는지다.",
    )

    doc.add_heading("1. 한눈에 보는 결론: 불균형보다 확률과 운영점이 더 중요했다", level=1)
    add_para(
        doc,
        "이 프로젝트는 HAM10000 피부 병변 데이터셋의 7-class 분류 문제에서 ConvNeXt-Tiny 기반 모델을 중심으로 성능 개선 과정을 추적한 실험이다. "
        "처음에는 극단적인 클래스 불균형이 가장 큰 병목이라고 보았다. 그래서 Focal Loss, WeightedRandomSampler, Class-Balanced Loss를 차례로 적용하였다. "
        "그러나 실험 결과 불균형 대응은 일부 소수 클래스 recall 개선에는 기여했지만, 최종 Macro-AUROC 관점에서는 ConvNeXt 기반 주 모델을 넘어서지 못했다.",
    )
    add_para(
        doc,
        "하지만 결과를 보면 문제는 단순히 '소수 클래스를 더 많이 보게 만들면 해결된다'에 가깝지 않았다. 이후 분석의 초점은 backbone, 입력 해상도, TTA, 확률 보정, 클래스별 threshold 조정으로 이동하였다. "
        "Baseline ConvNeXt 모델은 validation 기준 accuracy 0.8482, Macro-AUROC 0.966819를 기록했으며, flip TTA 적용 후 accuracy 0.8522, Macro-AUROC 0.9695로 개선되었다. "
        "Temperature Scaling은 accuracy를 유지하면서 NLL과 ECE를 낮춰 확률 신뢰도를 개선했고, class-wise threshold optimization은 melanoma, bcc, df에서 recall 또는 F1 개선을 확인하였다.",
    )
    add_para(
        doc,
        "따라서 최종 결론은 모델 구조 개선뿐 아니라 확률 보정과 클래스별 decision threshold 분석이 의료영상 분류 모델의 실질적 활용성을 높이는 데 중요하다는 것이다. "
        "단, 본 결과는 단일 validation split에 기반하므로 임상적 진단 성능을 단정하기보다 분류 보조 가능성을 확인한 결과로 해석해야 한다.",
    )

    doc.add_heading("2. 문제를 어떻게 다시 보게 되었는가", level=1)
    add_para(
        doc,
        "HAM10000은 피부 병변 dermoscopic 이미지 10,015장으로 구성된 공개 데이터셋이며, akiec, bcc, bkl, df, mel, nv, vasc의 7개 클래스를 포함한다. "
        "이 문제는 클래스 간 시각적 유사성이 높고 데이터 분포가 불균형하다는 점에서 일반적인 자연 이미지 분류보다 해석이 어렵다. 특히 melanoma와 nevus는 색소성 병변이라는 공통 특성 때문에 오분류가 발생하기 쉽다.",
    )
    add_para(
        doc,
        "의료영상 분류에서는 전체 accuracy만으로 모델의 실용성을 판단하기 어렵다. 다수 클래스인 nv가 전체 샘플의 큰 비중을 차지하기 때문에, 단순 accuracy는 소수 클래스 성능 저하를 가릴 수 있다. "
        "따라서 본 연구에서는 Macro-AUROC를 중심 지표로 사용하고, melanoma 등 주요 클래스의 precision, recall, F1 변화를 함께 검토하였다.",
    )
    add_para(
        doc,
        "본 보고서는 실험을 단순히 나열하지 않고, 초기 가설과 수정 가설의 흐름에 따라 구성하였다. 초반에는 클래스 불균형을 핵심 병목으로 보고 loss와 sampling 전략을 검증하였다. "
        "이후 결과를 바탕으로 확률 품질, calibration, 클래스별 threshold 설정이 실제 의사결정 단계에서 더 중요한 요소일 수 있다는 방향으로 분석을 확장하였다.",
    )
    add_para(
        doc,
        "자료의 기준도 구분해서 보았다. 기존 HAM10000_실험결과_리포트_개정판.pdf에는 강사님이 노트북에 적어 둔 '더 해볼 것' 항목, 즉 class weight, sampling, backbone, 해상도, TTA, 서브그룹 평가를 구현한 결과가 중심으로 담겨 있다. "
        "그 이후 추가된 Focal Loss, WeightedRandomSampler, Class-Balanced Loss, OVR Expert, Stacking, Temperature Scaling, Threshold Optimization은 최신 노트북 출력에서만 확인할 수 있으므로, 최종 판단은 최신 노트북 수치를 우선했다.",
    )

    doc.add_heading("3. 데이터와 평가 기준", level=1)
    add_para(
        doc,
        "데이터셋의 클래스 분포는 nv 6,705장, mel 1,113장, bkl 1,099장, bcc 514장, akiec 327장, vasc 142장, df 115장으로 확인된다. "
        "최다 클래스와 최소 클래스의 비율은 약 58:1이며, 이 불균형은 학습 중 다수 클래스 중심의 decision boundary를 만들 가능성이 있다.",
    )
    add_table(
        doc,
        "표 1. HAM10000 클래스 구성",
        ["클래스", "의미", "샘플 수", "해석상 주의점"],
        [
            ["nv", "Melanocytic nevi", "6,705", "다수 클래스이며 accuracy를 크게 좌우함"],
            ["mel", "Melanoma", "1,113", "FN 해석이 중요한 악성 병변 클래스"],
            ["bkl", "Benign keratosis-like lesions", "1,099", "mel/nv와 혼동 가능"],
            ["bcc", "Basal cell carcinoma", "514", "recall 개선 여부를 별도 확인"],
            ["akiec", "Actinic keratoses", "327", "샘플 수가 적어 변동성 주의"],
            ["vasc", "Vascular lesions", "142", "소수 클래스"],
            ["df", "Dermatofibroma", "115", "validation support가 매우 작음"],
        ],
        [1.8, 4.5, 1.7, 8.4],
    )
    add_para(
        doc,
        "Accuracy는 최종 예측 label이 정답과 일치하는 비율이므로, 다수 클래스 예측이 강한 모델에서 높게 나타날 수 있다. 반면 Macro-AUROC는 각 클래스를 one-vs-rest 문제로 보아 AUROC를 계산한 뒤 단순 평균하므로, 클래스별 ranking 품질을 더 균형 있게 반영한다. "
        "다만 AUROC는 threshold-independent 지표이기 때문에, threshold optimization으로 label decision rule을 바꾸더라도 AUROC 자체가 직접 개선되는 것은 아니다.",
    )

    doc.add_heading("4. 실험 전개: 처음 가설에서 최종 판단까지", level=1)
    add_para(
        doc,
        "전체 실험은 Baseline 설정, backbone 비교, 입력 해상도 및 TTA 개선, 클래스 불균형 대응, 전문가 모델 및 ensemble, calibration과 threshold 분석 순서로 진행하였다. "
        "각 단계에서는 이전 실험에서 드러난 한계를 다음 가설로 연결하였다.",
    )
    add_table(
        doc,
        "표 2. 실험별 목적 및 채택 여부",
        ["실험", "목적", "핵심 결과", "채택 판단"],
        [
            ["Baseline", "기준 성능과 오류 양상 확인", "ConvNeXt baseline Macro-AUROC 0.966819", "기준 모델"],
            ["Backbone 비교", "구조 차이에 따른 feature 품질 비교", "ConvNeXt-Tiny가 비교 backbone 중 우수", "채택"],
            ["224→384", "병변 경계와 세부 패턴 보존", "ResNet 기준 AUROC 0.899→0.907", "부분 채택"],
            ["Flip TTA", "추론 안정화", "AUROC 0.9695", "채택"],
            ["Focal Loss", "소수 클래스 강조", "AUROC 하락 및 accuracy 저하", "미채택"],
            ["WeightedRandomSampler", "소수 클래스 노출 증가", "df/akiec recall 개선, mel 제한", "보조 실험"],
            ["Class-Balanced Loss", "가중 과잉 완화", "Focal보다 안정적, 최종 미달", "대안 실험"],
            ["OVR Expert", "mel 보완", "mel AUROC/recall 일부 개선", "보완용"],
            ["Ensemble", "확률 결합 성능 확인", "Soft Voting/Ridge 차이 작음", "미채택"],
            ["Temperature Scaling", "확률 신뢰도 보정", "NLL/ECE 개선", "채택"],
            ["Threshold Optimization", "클래스별 sensitivity/F1 개선", "mel/bcc/df 개선 확인", "중요 분석"],
        ],
        [2.8, 5.0, 5.0, 2.5],
    )
    add_callout(
        doc,
        "초기 가설은 클래스 불균형이 핵심 병목이라는 것이었다. 그러나 불균형 대응 실험의 결과는 일부 recall 개선에 그쳤고, 최종 Macro-AUROC 개선은 제한적이었다. "
        "따라서 분석 가설은 확률 품질과 클래스별 decision threshold가 더 중요한 병목일 수 있다는 방향으로 수정되었다.",
    )

    doc.add_heading("4.1 Baseline 성능과 오류 해석", level=2)
    add_para(
        doc,
        "최종 분석의 기준이 되는 ConvNeXt baseline은 validation set에서 accuracy 0.8482, Macro-AUROC 0.966819를 기록하였다. "
        "표면적으로는 높은 AUROC를 보였지만 classification report를 보면 클래스별 recall 차이가 뚜렷하다. nv는 recall 0.952로 안정적이었으나, melanoma recall은 0.493으로 낮았고 df와 vasc도 각각 0.522, 0.607에 머물렀다.",
    )
    add_para(
        doc,
        "이 결과는 모델이 전체 ranking 품질은 높게 유지하면서도 argmax decision 단계에서는 일부 클래스의 sensitivity가 부족할 수 있음을 보여준다. "
        "따라서 본 연구에서는 AUROC와 accuracy뿐 아니라 class-wise recall, precision, F1, 그리고 threshold 조정 후의 trade-off를 함께 분석하였다.",
    )
    add_table(
        doc,
        "표 3. ConvNeXt baseline 클래스별 성능",
        ["클래스", "AUROC", "Precision", "Recall", "F1", "Support", "해석"],
        [
            ["akiec", "0.9635", "0.623", "0.662", "0.642", "65", "소수 클래스이나 AUROC는 높음"],
            ["bcc", "0.9830", "0.689", "0.816", "0.747", "103", "recall은 상대적으로 양호"],
            ["bkl", "0.9477", "0.693", "0.709", "0.701", "220", "중간 수준의 균형"],
            ["df", "0.9906", "0.750", "0.522", "0.615", "23", "support가 작아 변동성 큼"],
            ["mel", "0.9230", "0.748", "0.493", "0.595", "223", "FN 감소가 핵심 과제"],
            ["nv", "0.9615", "0.908", "0.952", "0.929", "1,341", "다수 클래스, 안정적"],
            ["vasc", "0.9984", "1.000", "0.607", "0.756", "28", "precision 높지만 recall 제한"],
        ],
        [1.6, 1.8, 1.8, 1.6, 1.5, 1.7, 6.4],
    )

    doc.add_heading("5. 첫 번째 전환점: Backbone과 입력 품질", level=1)
    add_para(
        doc,
        "Backbone 비교는 동일한 데이터 split과 classifier 학습 전략을 기준으로 ResNet18, EfficientNet-B0, ConvNeXt-Tiny를 비교하였다. "
        "ConvNeXt-Tiny는 validation Macro-AUROC 0.940512로 ResNet18과 EfficientNet-B0보다 높은 값을 보였다. 이 결과는 ConvNeXt 계열의 modern convolution 구조가 피부 병변 이미지의 texture와 local pattern 표현에 더 유리했을 가능성을 시사한다.",
    )
    add_table(
        doc,
        "표 4. Backbone 비교",
        ["Backbone", "전략", "Best val Macro-AUROC", "Best val Accuracy", "해석"],
        [
            ["ResNet18", "classifier", "0.898899", "0.735", "가볍지만 표현력 제한"],
            ["EfficientNet-B0", "classifier", "0.894029", "0.705", "효율적이나 본 실험에서는 낮음"],
            ["ConvNeXt-Tiny", "classifier", "0.940512", "0.782", "비교군 중 가장 우수"],
        ],
        [3.0, 2.2, 3.0, 2.8, 5.0],
    )
    add_chart(doc, charts["backbone"], "그림 1. Backbone별 Macro-AUROC 비교. ConvNeXt-Tiny가 가장 높은 검증 AUROC를 보였다.")
    add_para(
        doc,
        "입력 해상도 실험에서는 224에서 384로 크기를 늘렸을 때 ResNet18 기준 best validation Macro-AUROC가 0.899467에서 0.906514로 상승하였다. "
        "의료영상에서는 병변의 경계, 색 변화, 국소 texture가 분류 근거로 작용할 수 있으므로 해상도 증가는 세부 패턴 보존에 의미가 있다. 다만 계산 시간이 증가했기 때문에 성능 향상과 비용 사이의 trade-off를 함께 고려해야 한다.",
    )
    add_table(
        doc,
        "표 5. 입력 해상도 및 TTA 비교",
        ["실험", "Accuracy", "Macro-AUROC", "주요 해석"],
        [
            ["IMG_SIZE 224", "0.735", "0.899467", "기본 해상도 기준"],
            ["IMG_SIZE 384", "0.757", "0.906514", "경계와 세부 패턴 보존 효과 가능"],
            ["ConvNeXt baseline", "0.8482", "0.966819", "최종 분석의 기준 확률"],
            ["ConvNeXt + flip TTA", "0.8522", "0.9695", "추론 안정화와 소폭 성능 개선"],
        ],
        [3.5, 2.4, 2.8, 7.0],
    )
    add_para(
        doc,
        "Flip TTA는 원본과 좌우 반전 이미지를 함께 추론한 뒤 평균 확률을 사용하는 방식이다. 피부 병변 분류에서는 좌우 방향 자체보다 병변 내부 패턴이 더 중요하므로, flip TTA는 예측 분산을 줄이는 후처리로 해석할 수 있다. "
        "본 실험에서 TTA는 accuracy와 Macro-AUROC를 모두 개선했기 때문에 최종 모델의 추론 안정화 기법으로 채택하였다.",
    )
    add_para(
        doc,
        "다만 TTA는 모델이 학습하지 못한 새로운 정보를 추가하는 기법이 아니다. 같은 이미지의 변환된 관찰값을 평균하여 확률 변동을 줄이는 방식이므로, 성능 향상 폭은 일반적으로 제한적이다. "
        "본 실험에서도 개선 폭은 크지 않았지만 방향성이 일관되었고, 별도의 재학습 없이 적용할 수 있다는 점에서 최종 추론 절차에 포함할 근거가 충분했다.",
    )

    doc.add_heading("6. 두 번째 전환점: 불균형 대응만으로는 충분하지 않았다", level=1)
    add_para(
        doc,
        "클래스 불균형 대응 실험은 초기 가설을 직접 검증하기 위한 단계였다. Focal Loss, WeightedRandomSampler, Class-Balanced Loss는 모두 소수 클래스의 학습 기여를 키우는 목적을 가진다. "
        "그러나 세 실험 모두 최종 ConvNeXt baseline 또는 TTA 결과를 넘어서지 못했다.",
    )
    add_para(
        doc,
        "Focal Loss + class weight는 df, vasc, akiec 등 소수 클래스에 큰 가중치를 부여하였다. 그 결과 일부 소수 클래스 recall은 증가했으나, 전체 accuracy가 0.555 수준으로 낮아지고 Macro-AUROC도 약 0.9098에 머물렀다. "
        "이는 소수 클래스 가중이 과도하게 작동하여 다수 클래스인 nv의 예측 안정성을 훼손했을 가능성을 보여준다.",
    )
    add_para(
        doc,
        "WeightedRandomSampler + CrossEntropy는 학습 중 소수 클래스 노출 빈도를 높이는 방식이다. 이 실험은 akiec, df 등 일부 클래스 recall을 개선했지만 melanoma recall은 0.363 수준으로 baseline보다 낮았다. "
        "따라서 melanoma 성능 병목은 단순한 샘플 수 부족만으로 설명하기 어렵고, mel과 nv/bkl 사이의 시각적 유사성 및 확률 분포 문제가 함께 작용한 것으로 해석된다.",
    )
    add_para(
        doc,
        "Class-Balanced Loss는 Effective Number of Samples 기반으로 class weight를 완화해 Focal Loss보다 안정적인 결과를 보였다. melanoma recall은 0.556으로 baseline보다 개선되었으나 Macro-AUROC는 0.9457로 최종 ConvNeXt baseline에는 미치지 못했다. "
        "따라서 이 방법은 recall 보완 목적의 대안 실험으로 해석하고 최종 모델로는 채택하지 않았다.",
    )
    add_table(
        doc,
        "표 6. 불균형 대응 실험 비교",
        ["실험", "Accuracy", "Macro-AUROC", "mel recall", "채택 여부 및 근거"],
        [
            ["ConvNeXt baseline", "0.8482", "0.966819", "0.4933", "최종 기준"],
            ["Focal Loss + class weight", "0.555", "0.9098", "0.605", "가중 과잉 가능성, 미채택"],
            ["Focal + TTA", "0.5612", "0.9111", "0.605", "TTA 후에도 최종 미달"],
            ["WeightedRandomSampler + CE", "0.784", "0.9505", "0.363", "일부 recall 개선, mel 제한"],
            ["Class-Balanced Loss", "0.756", "0.9457", "0.556", "Focal보다 안정적이나 최종 미달"],
            ["Class-Balanced + TTA", "0.7629", "0.9457", "0.565", "recall 보완용 대안"],
        ],
        [3.7, 2.2, 2.6, 2.3, 6.0],
    )
    add_chart(doc, charts["imbalance"], "그림 2. 불균형 대응 실험 비교. 소수 클래스 보정은 일부 recall에는 효과가 있었지만 최종 Macro-AUROC 기준 baseline을 넘지 못했다.")

    doc.add_heading("7. 세 번째 전환점: Expert와 Ensemble은 보완책이었다", level=1)
    add_para(
        doc,
        "OVR Expert는 melanoma와 나머지 클래스를 구분하는 이진 전문가 모델로 설계되었다. 이 접근은 전체 7-class decision boundary를 직접 바꾸기보다, melanoma 확률을 별도로 보완하는 목적을 가진다. "
        "노트북의 OVR 실험에서는 OVR Expert가 melanoma AUROC와 recall 개선 가능성을 보였으나, 최종 모델을 대체하기보다는 melanoma 보완용 실험으로 해석하는 것이 타당하다.",
    )
    add_para(
        doc,
        "OVR Soft Ensemble은 ConvNeXt의 기존 mel 확률과 OVR Expert의 mel 확률을 alpha 비율로 soft blending하는 방식이다. alpha=0.05에서는 Macro-AUROC가 0.966819에서 0.966891로 매우 미세하게 개선되었다. "
        "그러나 alpha가 커질수록 mel recall은 증가하는 반면 Macro-AUROC는 감소했다. 이는 ConvNeXt의 기존 7-class 확률 구조가 이미 안정적이며, 외부 mel 확률을 강하게 주입하면 전체 ranking 품질이 손상될 수 있음을 의미한다.",
    )
    add_table(
        doc,
        "표 7. OVR Soft Ensemble alpha sweep 요약",
        ["alpha", "Accuracy", "Macro-AUROC", "mel AUROC", "mel Precision", "mel Recall", "해석"],
        [
            ["0.00", "0.8482", "0.966819", "0.922971", "0.7483", "0.4933", "기준 ConvNeXt 확률"],
            ["0.05", "0.8482", "0.966891", "0.924034", "0.7383", "0.4933", "AUROC 미세 개선"],
            ["0.35", "0.8512", "0.963875", "0.908031", "0.7229", "0.5381", "recall 증가, AUROC 감소"],
            ["0.65", "0.8477", "0.961526", "0.896468", "0.6279", "0.6054", "mel recall은 증가하나 전체 ranking 손상"],
            ["0.80", "0.8427", "0.960279", "0.890074", "0.5644", "0.6682", "강한 blending은 미채택"],
        ],
        [1.5, 2.1, 2.5, 2.3, 2.5, 2.1, 5.4],
    )
    add_chart(doc, charts["ovr"], "그림 3. OVR alpha trade-off. alpha가 커지면 mel recall은 증가하지만 Macro-AUROC는 감소했다.")
    add_para(
        doc,
        "alpha sweep은 threshold optimization과 다른 방식의 trade-off를 보여준다. OVR blending은 확률 자체를 바꾸기 때문에 AUROC와 class-wise AUROC에 영향을 줄 수 있다. "
        "반면 threshold optimization은 이미 계산된 확률의 decision rule만 바꾸므로 AUROC reference는 동일하게 유지된다. 이 차이를 구분하는 것이 결과 해석에서 중요하다.",
    )
    add_para(
        doc,
        "Ensemble 실험에서는 ResNet18, EfficientNet-B0, ConvNeXt-Tiny의 확률을 결합하였다. Soft Voting은 Macro-AUROC 0.9472를 기록했으며, Logistic Stacking은 meta-valid split에서 accuracy 0.8114로 높았지만 Macro-AUROC는 0.9431로 낮았다. "
        "Ridge Stacking은 Macro-AUROC 0.9449로 Stacking 계열에서 더 안정적이었으나, Soft Voting과 차이가 작고 최종 ConvNeXt baseline에는 미치지 못했다.",
    )
    add_table(
        doc,
        "표 8. Ensemble 비교",
        ["방법", "평가 split", "Accuracy", "Macro-AUROC", "해석"],
        [
            ["Single ResNet18", "val", "0.7339", "0.8989", "단일 약한 기준"],
            ["Single EfficientNet-B0", "val", "0.6965", "0.8940", "본 실험에서 낮음"],
            ["Single ConvNeXt-Tiny", "val", "0.7718", "0.9405", "backbone 비교 기준"],
            ["Soft Voting", "val", "0.7993", "0.9472", "확률 평균으로 개선"],
            ["Logistic Stacking", "meta-valid", "0.8114", "0.9431", "accuracy는 높지만 AUROC 낮음"],
            ["Ridge Stacking", "meta-valid", "0.8044", "0.9449", "Stacking 계열 중 우수"],
            ["Weighted Ensemble", "meta-valid", "0.7974", "0.9452", "기대만큼 향상되지 않음"],
        ],
        [3.3, 2.6, 2.1, 2.6, 5.2],
    )
    add_callout(
        doc,
        "Ensemble 결과는 모델을 많이 결합한다고 항상 최종 지표가 개선되는 것은 아니라는 점을 보여준다. 특히 meta-valid split 기반 Stacking 결과는 split 차이의 영향을 받으므로, Soft Voting과 Ridge의 우위를 통계적으로 단정하지 않았다.",
    )

    doc.add_heading("8. 마지막 전환점: 확률 보정과 클래스별 threshold", level=1)
    add_para(
        doc,
        "Temperature Scaling은 모델의 logits를 하나의 온도 파라미터 T로 나누어 softmax 확률의 sharpness를 조절하는 calibration 기법이다. 이 방법은 argmax label을 직접 바꾸는 것이 아니라, 예측 확률이 실제 정답 가능성과 더 잘 대응하도록 후처리하는 기법이다.",
    )
    add_para(
        doc,
        "본 실험에서 최적 temperature는 약 1.354742였다. Accuracy는 0.848228로 유지되었고, Macro-AUROC는 0.966819에서 0.966914로 소폭 증가하였다. 더 중요한 변화는 NLL이 0.464353에서 0.439376으로 감소하고, ECE가 0.044550에서 0.019226으로 크게 감소한 점이다. "
        "의료 AI에서는 단순한 label 정확도뿐 아니라 예측 확률의 신뢰도가 중요하므로, calibration 개선은 모델의 분류 보조 가능성을 평가하는 데 의미가 있다.",
    )
    add_table(
        doc,
        "표 9. Temperature Scaling 전후 비교",
        ["설정", "Temperature", "Accuracy", "Macro-AUROC", "NLL", "ECE"],
        [
            ["Before", "1.000000", "0.848228", "0.966819", "0.464353", "0.044550"],
            ["After", "1.354742", "0.848228", "0.966914", "0.439376", "0.019226"],
        ],
        [2.7, 2.7, 2.5, 3.0, 2.4, 2.4],
    )
    add_chart(doc, charts["calibration"], "그림 4. Temperature Scaling 전후 NLL/ECE 변화. label은 유지하면서 확률 신뢰도 지표가 개선되었다.")
    add_para(
        doc,
        "Class-wise Threshold Optimization은 모든 클래스에 동일한 argmax decision rule을 적용하는 것이 항상 최적인지 검토하기 위한 실험이다. "
        "본 실험에서는 melanoma, bcc, df를 대상으로 threshold를 조정했고, 각 클래스의 recall, precision, F1, accuracy 변화를 비교하였다.",
    )
    add_table(
        doc,
        "표 10. Temperature Scaling class-wise AUROC 변화",
        ["클래스", "Before AUROC", "After AUROC", "Delta", "해석"],
        [
            ["df", "0.990580", "0.991546", "+0.000966", "가장 큰 증가이나 support 작음"],
            ["bcc", "0.982999", "0.983500", "+0.000501", "소폭 개선"],
            ["bkl", "0.947688", "0.947999", "+0.000311", "소폭 개선"],
            ["nv", "0.961544", "0.961778", "+0.000234", "거의 유지"],
            ["akiec", "0.963507", "0.963531", "+0.000024", "변화 미미"],
            ["vasc", "0.998445", "0.998246", "-0.000199", "미세 하락"],
            ["mel", "0.922971", "0.921797", "-0.001174", "mel AUROC는 소폭 하락"],
        ],
        [1.8, 2.7, 2.7, 2.0, 6.2],
    )
    add_para(
        doc,
        "Temperature Scaling 후 class-wise AUROC 변화는 매우 작았다. 따라서 이 실험의 핵심은 AUROC를 크게 높인 것이 아니라 NLL과 ECE를 낮춰 확률 신뢰도를 개선했다는 점이다. "
        "특히 mel AUROC가 소폭 하락했음에도 전체 calibration 지표가 개선되었으므로, 보고서에서는 이를 성능 향상 기법이라기보다 확률 품질 보정 기법으로 설명해야 한다.",
    )
    add_table(
        doc,
        "표 11. Threshold Optimization 전후 비교",
        ["클래스", "Threshold", "Accuracy", "Precision", "Recall", "F1", "해석"],
        [
            ["mel", "0.28", "0.8482\n→ 0.8447", "0.7483\n→ 0.6505", "0.4933\n→ 0.6009", "0.5946\n→ 0.6247", "recall 크게 증가, accuracy 소폭 감소"],
            ["bcc", "0.37", "0.8482\n→ 0.8497", "0.6885\n→ 0.6739", "0.8155\n→ 0.9029", "0.7467\n→ 0.7718", "recall/F1 증가, accuracy 소폭 증가"],
            ["df", "0.18", "0.8482\n→ 0.8497", "0.7500\n→ 0.7895", "0.5217\n→ 0.6522", "0.6154\n→ 0.7143", "precision/recall/F1/accuracy 모두 개선"],
        ],
        [1.5, 1.8, 2.8, 2.8, 2.8, 2.8, 4.5],
    )
    doc.add_page_break()
    add_table(
        doc,
        "표 12. Threshold Optimization 변화량 요약",
        ["Class", "Threshold", "Accuracy 변화", "Recall 변화", "Precision 변화", "F1 변화", "핵심 해석"],
        [
            ["mel", "0.28", "-0.0035", "+0.1076", "-0.0978", "+0.0301", "FP 허용, sensitivity 개선"],
            ["bcc", "0.37", "+0.0015", "+0.0874", "-0.0146", "+0.0251", "recall/F1/accuracy 동시 개선"],
            ["df", "0.18", "+0.0015", "+0.1304", "+0.0395", "+0.0989", "모든 지표 개선"],
        ],
        [1.4, 1.7, 2.0, 1.9, 2.1, 1.7, 5.9],
    )
    add_para(
        doc,
        "이 변화량 표는 본 보고서에서 가장 중요한 해석 근거다. mel은 threshold를 낮추면서 recall을 약 10.8%p 높였고, accuracy 감소는 약 0.35%p로 제한되었다. "
        "bcc는 recall과 F1뿐 아니라 accuracy도 소폭 개선되어 기존 판단 기준이 bcc에 대해 다소 보수적이었음을 보여준다. "
        "df는 precision, recall, F1, accuracy가 동시에 개선되어 단일 argmax 규칙보다 클래스별 operating point가 더 적절할 수 있음을 가장 분명하게 보여준다.",
    )
    doc.add_page_break()
    add_chart(doc, charts["threshold"], "그림 5. Threshold 조정 전후 recall/F1 변화. mel, bcc, df 모두 sensitivity 또는 F1 관점에서 개선이 확인되었다.")
    add_para(
        doc,
        "melanoma의 경우 threshold 0.28에서 recall이 0.4933에서 0.6009로 증가했으며, accuracy 감소는 0.0035 수준으로 제한적이었다. precision은 감소했으므로 FP 증가 가능성을 동반하지만, 의료영상 분류에서는 melanoma FN을 줄이는 것이 중요한 운영 목표가 될 수 있다. "
        "즉 false positive를 일부 허용하면서 민감도를 높이는 전형적인 의료 screening 모델의 특성을 보였다.",
    )
    add_para(
        doc,
        "bcc는 threshold 0.37에서 recall이 0.8155에서 0.9029로 상승했고, F1도 0.7467에서 0.7718로 증가했다. 보통 recall을 높이면 accuracy가 떨어지는 경우가 많지만, 이 실험에서는 accuracy도 0.8482에서 0.8497로 소폭 증가했다. "
        "이는 ConvNeXt가 bcc에 대해 다소 보수적인 threshold를 사용하고 있었고, threshold 조정만으로 놓치던 bcc 샘플을 회수할 수 있었음을 시사한다.",
    )
    add_para(
        doc,
        "가장 인상적인 결과는 df였다. threshold 0.18에서 precision, recall, F1, accuracy가 모두 개선되었다. 일반적으로 precision과 recall은 trade-off 관계에 있지만, df에서는 둘 다 좋아졌다. "
        "이는 현재 ConvNeXt가 df에 대해 threshold를 지나치게 높게 설정하고 있었고, 클래스별 operating point를 조정하는 것만으로 단일 argmax 규칙보다 더 적절한 의사결정이 가능했음을 보여준다.",
    )
    add_para(
        doc,
        "단, threshold 선택은 validation set의 확률 분포에 의존한다. 따라서 실제 배포에서는 독립 test set 또는 시간적으로 분리된 external validation에서 동일 threshold가 유지되는지 확인해야 한다. "
        "또한 threshold를 낮추면 일반적으로 FP가 증가하므로, 의료 현장에서는 추가 검사 비용과 FN 감소의 이득을 함께 평가해야 한다.",
    )
    add_callout(
        doc,
        "Threshold Optimization의 결론은 모든 클래스에 동일한 argmax decision rule을 적용하는 것이 최적은 아니라는 점이다. 클래스별 확률 분포와 오류 비용을 고려한 threshold 설정은 실제 배포 환경에서 sensitivity와 F1을 조정하는 효과적인 방법이 될 수 있다.",
    )

    doc.add_heading("9. 종합 비교: 최종 판단 기준", level=1)
    add_para(
        doc,
        "전체 실험을 종합하면, 본 프로젝트에서 가장 강한 정량 성능은 ConvNeXt 기반 baseline과 TTA 조합에서 나타났다. 불균형 대응 기법은 소수 클래스 recall 개선에는 도움이 되었으나, Macro-AUROC 기준 최종 모델을 대체할 만큼의 개선을 만들지는 못했다. "
        "이는 HAM10000에서 성능 병목이 단순한 클래스 빈도 문제가 아니라, 클래스 간 시각적 유사성, 확률 분포, decision threshold 문제와 결합되어 있음을 시사한다.",
    )
    add_table(
        doc,
        "표 13. 전체 실험 요약표",
        ["범주", "대표 실험", "핵심 수치", "채택/미채택 근거"],
        [
            ["기준 모델", "ConvNeXt baseline", "Acc 0.8482 / AUROC 0.966819", "최종 분석의 중심 모델"],
            ["추론 개선", "Flip TTA", "Acc 0.8522 / AUROC 0.9695", "안정적 개선으로 채택"],
            ["해상도", "224→384", "AUROC 0.8995→0.9065", "세부 패턴 보존 가능성"],
            ["불균형", "Focal Loss", "AUROC 0.9098", "과도한 가중으로 미채택"],
            ["불균형", "WeightedRandomSampler", "AUROC 0.9505", "mel 개선 제한"],
            ["불균형", "Class-Balanced Loss", "AUROC 0.9457", "대안 실험으로 해석"],
            ["Expert", "OVR Soft alpha=0.05", "AUROC 0.966891", "개선 폭 매우 작음"],
            ["Ensemble", "Soft Voting/Ridge", "0.9472 / 0.9449", "최종 baseline 미달"],
            ["Calibration", "Temperature Scaling", "ECE 0.04455→0.019226", "확률 신뢰도 개선"],
            ["Threshold", "mel/bcc/df", "recall 및 F1 개선", "운영점 분석으로 중요"],
        ],
        [2.3, 3.5, 4.3, 6.0],
    )
    add_para(
        doc,
        "초기에는 클래스 불균형이 핵심 병목일 것으로 예상했으나, 실험 결과는 더 복합적이었다. Focal Loss는 소수 클래스에 대한 민감도를 높였지만 전체 성능을 크게 떨어뜨렸고, WeightedRandomSampler는 일부 소수 클래스에 효과가 있었지만 melanoma 성능을 충분히 개선하지 못했다. "
        "Class-Balanced Loss는 상대적으로 안정적이었으나 최종 Macro-AUROC 기준으로는 ConvNeXt baseline보다 낮았다.",
    )
    add_para(
        doc,
        "반면 Temperature Scaling과 Threshold Optimization은 모델 구조를 바꾸지 않으면서도 의료영상 분류에서 중요한 확률 신뢰도와 sensitivity trade-off를 조정했다. "
        "Temperature Scaling은 label 예측을 유지하면서 NLL과 ECE를 낮추었고, threshold 조정은 melanoma, bcc, df에서 recall 또는 F1을 개선했다. 따라서 최종 보고서의 핵심 메시지는 모델 자체의 성능뿐 아니라 확률 보정과 클래스별 운영점 분석이 필요하다는 것이다.",
    )
    doc.add_page_break()
    add_table(
        doc,
        "표 14. 보고서 가치가 높은 실험 순위",
        ["순위", "실험", "보고서에서 중요한 이유"],
        [
            ["1", "Class-wise Threshold Optimization", "mel/bcc/df 비교로 동일 argmax 규칙의 한계를 직접 확인"],
            ["2", "OVR Expert mel vs rest", "melanoma 보완 가능성을 별도 전문가 모델로 검증"],
            ["3", "Ridge Stacking", "Stacking 계열 중 가장 안정적이나 Soft Voting과 차이는 작음"],
            ["4", "Resolution 224→384", "세부 texture와 병변 경계 보존의 효과를 확인"],
            ["5", "TTA", "재학습 없이 추론 안정성을 높임"],
            ["6", "WeightedRandomSampler", "일부 recall 개선, mel 병목 개선은 제한적"],
            ["7", "Class-Balanced Loss", "Focal보다 안정적이나 최종 모델에는 미달"],
            ["8", "Soft Voting", "단순 확률 평균의 기준점과 한계 확인"],
            ["9", "Logistic Stacking", "accuracy는 높였지만 Macro-AUROC는 낮아짐"],
            ["10", "Focal Loss", "과도한 소수 클래스 가중의 위험 확인"],
        ],
        [1.2, 4.4, 10.8],
        body_size=7.35,
        header_size=7.8,
        margin_y=45,
        margin_x=90,
    )
    add_callout(
        doc,
        "최종적으로 가장 중요한 실험은 Class-wise Threshold Optimization이다. 이 실험은 'threshold를 낮추면 recall이 올라간다'는 일반론을 넘어서, 클래스마다 최적 threshold와 성능 변화 양상이 다르다는 점을 보여주었다. "
        "특히 bcc와 df는 threshold 조정만으로 recall과 F1이 동시에 개선되었고, df는 precision과 accuracy까지 함께 개선되었다.",
    )

    doc.add_page_break()
    doc.add_heading("10. 결론 및 향후 연구", level=1)
    add_para(
        doc,
        "본 연구에서는 HAM10000 피부 병변 분류 문제에 대해 ConvNeXt-Tiny 기반 모델을 중심으로 다양한 성능 향상 실험을 수행하였다. "
        "실험 결과, 단순한 손실 함수 변경이나 샘플링 전략보다 입력 해상도, TTA, 확률 보정, 클래스별 threshold 분석이 더 실질적인 효과를 보였다.",
    )
    add_para(
        doc,
        "특히 Temperature Scaling은 모델의 예측 확률 신뢰도를 개선했으며, Class-wise Threshold Optimization은 melanoma, bcc, df에서 sensitivity 및 F1 개선을 보였다. "
        "따라서 최종적으로 ConvNeXt-Tiny 기반 모델을 중심으로 하되, 확률 보정과 클래스별 threshold 분석을 결합하는 방향이 가장 타당한 결론이다.",
    )
    add_para(
        doc,
        "본 연구의 한계는 단일 validation split에 기반하고, 외부 검증 데이터셋이나 임상 메타데이터 결합을 포함하지 않았다는 점이다. 또한 threshold는 validation 분포에 민감할 수 있으므로 실제 적용 전에는 독립 test set과 교차 검증이 필요하다.",
    )
    add_para(doc, "향후 연구 방향은 다음과 같다.")
    add_bullets(
        doc,
        [
            "ConvNeXt V2, Vision Transformer 등 더 다양한 backbone 비교",
            "나이, 성별, 병변 위치 등 임상 메타데이터와 이미지 feature의 결합",
            "Cost-sensitive learning과 클래스별 threshold를 함께 고려한 의사결정 분석",
            "melanoma 중심 다단계 Expert 모델 구성",
            "Temperature Scaling 외 Platt Scaling, Isotonic Regression 등 calibration 기법 비교",
            "외부 데이터셋 기반 generalization 평가",
        ],
    )

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_report()
    print(OUT_DOCX)
